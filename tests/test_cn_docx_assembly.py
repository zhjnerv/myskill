"""中国发明专利申请 DOCX 组装合同回归测试。"""

from __future__ import annotations

import importlib.util
import os
import shutil
import sys
from pathlib import Path
from zipfile import ZipFile

import pytest
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
SCRIPT = (
    ROOT
    / "skills"
    / "cn-patent-application-creator"
    / "scripts"
    / "assemble_application_docx.py"
)
REFERENCE = (
    ROOT
    / "skills"
    / "cn-patent-application-creator"
    / "references"
    / "docx-assembly.md"
)


def _load_module():
    spec = importlib.util.spec_from_file_location("cn_docx_assembler", SCRIPT)
    assert spec is not None and spec.loader is not None
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    return module


ASSEMBLER = _load_module()


def _write_test_template(path: Path) -> None:
    document = Document()
    for name in ("Normal (Web)", "正文2", "附图图号"):
        if name not in document.styles:
            document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    paragraph = document.add_paragraph()
    paragraph.add_run("模板权利要求")
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), "1")
    numpr.extend((ilvl, numid))
    ppr.append(numpr)
    headers = ["权利要求书", "说明书", "说明书附图", "说明书摘要", "摘要附图"]
    document.sections[0].header.paragraphs[0].text = headers[0]
    for header in headers[1:]:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        section.header.is_linked_to_previous = False
        section.header.paragraphs[0].text = header
    document.save(path)


def test_docx_assembly_skill_assets_are_utf8_and_documented():
    for path in (SCRIPT, REFERENCE):
        raw = path.read_bytes()
        assert raw
        assert not raw.startswith(b"\xef\xbb\xbf")
        raw.decode("utf-8")

    creator = ROOT / "skills" / "cn-patent-application-creator"
    skill = (creator / "SKILL.md").read_text(encoding="utf-8")
    docx_reference = (creator / "references" / "docx-assembly.md").read_text(encoding="utf-8")
    # SKILL.md 只保留入口指针，DOCX 技术细节按阶段加载到 reference。
    for required in (
        "assemble_application_docx.py",
        "references/docx-assembly.md",
    ):
        assert required in skill
    combined = skill + "\n" + docx_reference
    for required in (
        "m:oMath",
        "Strong",
        "visual_review_completed",
        "视觉检查默认关闭",
    ):
        assert required in combined
    agents = (ROOT / "AGENTS.md").read_text(encoding="utf-8")
    readme = (ROOT / "README.md").read_text(encoding="utf-8")
    assert "只有用户明确要求视觉检查时才导出 PDF/PNG" in agents
    assert "不提供业务型 Slash Command" in readme
    assert "python-docx" in readme


def test_formula_parser_supports_explicit_and_legacy_blocks(tmp_path):
    module = ASSEMBLER
    spec = tmp_path / "说明书.md"
    spec.write_text(
        """# 示例装置

## 技术领域

患者年龄为 $a$，转换系数为 $C(P,a)$。

## 背景技术

$$
C(P,a)=C(P,a_i)+(a-a_i)×[C(P,a_(i+1))-C(P,a_i)]/(a_(i+1)-a_i)
$$

## 发明内容

```math
ED_CT=Q_CT×C_CT(P,a)
```

## 附图说明

图1是示意图。

## 具体实施方式

m_k=m_(k-1)+(x_k-m_(k-1))/k。
""",
        encoding="utf-8",
    )

    items = module.parse_specification(spec)
    assert [item.kind for item in items].count("formula") == 3
    assert [item.text for item in items if item.kind == "heading"] == [
        "技术领域",
        "背景技术",
        "发明内容",
        "附图说明",
        "具体实施方式",
    ]
    symbols = module.collect_math_symbols(items)
    tokens = module.tokenize_inline_math("患者年龄为 $a$，转换系数为 $C(P,a)$。", symbols)
    assert [value for is_math, value in tokens if is_math] == ["a", "C(P,a)"]


def test_claim_abstract_and_figure_contracts(tmp_path):
    module = ASSEMBLER
    claims = tmp_path / "权利要求书.md"
    claims.write_text("# 权利要求书\n\n1. 第一项。\n\n2. 第二项。\n", encoding="utf-8")
    assert module.parse_claims(claims) == ["第一项。", "第二项。"]

    compact_claims = tmp_path / "紧凑权利要求书.md"
    compact_claims.write_text(
        "# 权利要求书\n\n1. 第一项第一行。\n第二行。\n2. 第二项。\n3. 第三项。\n",
        encoding="utf-8",
    )
    assert module.parse_claims(compact_claims) == [
        "第一项第一行。 第二行。",
        "第二项。",
        "第三项。",
    ]

    abstract = tmp_path / "说明书摘要.md"
    abstract.write_text("# 示例装置\n\n这是摘要。\n\n摘要附图：图2。\n", encoding="utf-8")
    assert module.parse_abstract(abstract) == ("这是摘要。", 2)

    drawing_dir = tmp_path / "说明书附图"
    drawing_dir.mkdir()
    for number in (1, 2):
        (drawing_dir / f"图{number}.png").write_bytes(b"png-placeholder")
    index = tmp_path / "说明书附图.md"
    index.write_text(
        """# 说明书附图

## 图1 系统结构示意图

![图1](说明书附图/图1.png)

## 图2 方法流程图

![图2](说明书附图/图2.png)
""",
        encoding="utf-8",
    )
    figures = module.parse_figures(index)
    assert [(figure.number, figure.title) for figure in figures] == [
        (1, "系统结构示意图"),
        (2, "方法流程图"),
    ]



def test_specification_contract_strips_paragraph_numbers_and_requires_concise_figures(tmp_path):
    module = ASSEMBLER
    spec = tmp_path / "说明书.md"
    spec.write_text(
        "# 示例装置\n\n## 技术领域\n\n[0001] 涉及示例技术。\n\n"
        "## 背景技术\n\n现有方案不足。\n\n## 发明内容\n\n提供示例方案。\n\n"
        "## 附图说明\n\n图1为本申请一实施方式中的结构图。\n\n图中：100-示例模块。\n\n"
        "## 具体实施方式\n\n以下结合附图说明本发明的具体实施例。\n\n### 实施例1。\n\n"
        "如图1所示，示例模块100执行处理。\n",
        encoding="utf-8",
    )
    items = module.parse_specification(spec)
    assert all(not item.text.startswith("[0001]") for item in items)
    figures = [module.FigureSpec(1, "结构图", tmp_path / "图1.png")]
    module.validate_specification_structure(items, figures)

    bad = [module.SpecItem(item.kind, item.text) for item in items]
    figure_index = next(i for i, item in enumerate(bad) if item.text.startswith("图1为"))
    bad[figure_index] = module.SpecItem("body", "图1示出了示例模块的内部连接、数据来源和完整处理过程。")
    with pytest.raises(ValueError, match="一图一句"):
        module.validate_specification_structure(bad, figures)


def test_method_claim_steps_are_split_into_template_paragraphs():
    module = ASSEMBLER
    parts = module.split_claim_paragraphs(
        "一种方法，其特征在于，包括：S101：执行第一步骤；S102：执行第二步骤。"
    )
    assert parts == [
        "一种方法，其特征在于，包括：",
        "S101：执行第一步骤；",
        "S102：执行第二步骤。",
    ]

def test_direct_omml_builds_native_fraction_subscript_and_superscript(tmp_path):
    module = ASSEMBLER
    formula = "C(P,a_i)+(x^2)/(y_1-z)"
    root = module.build_omath(formula)
    xml = etree.tostring(root)
    namespaces = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
    parsed = etree.fromstring(xml)
    assert len(parsed.xpath(".//m:f", namespaces=namespaces)) == 1
    assert len(parsed.xpath(".//m:sSub", namespaces=namespaces)) == 2
    assert len(parsed.xpath(".//m:sSup", namespaces=namespaces)) == 1

    document = Document()
    paragraph = document.add_paragraph("公式：")
    registry = module.MathRegistry()
    module.append_omath(paragraph, registry, formula)
    output = tmp_path / "native-math.docx"
    document.save(output)
    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
    package_root = etree.fromstring(document_xml)
    assert len(package_root.xpath(".//m:oMath", namespaces=namespaces)) == 1
    assert b"[[EQ" not in document_xml


def test_linux_native_formula_no_longer_requires_word(monkeypatch, tmp_path):
    module = ASSEMBLER
    monkeypatch.setattr(module.sys, "platform", "linux")
    document = Document()
    paragraph = document.add_paragraph()
    registry = module.MathRegistry()
    module.append_omath(paragraph, registry, "ED_CT=Q_CT×C_CT(P,a)")
    output = tmp_path / "native-math.docx"
    document.save(output)
    assert module.process_with_word(output, registry, None) is None
    package = module.inspect_package(output, 1, "Heading1", "Strong")
    assert package["math_count"] == 1


@pytest.mark.skipif(
    os.environ.get("CN_PATENT_RUN_VISUAL_TESTS") != "1"
    or shutil.which("libreoffice") is None or shutil.which("pdfinfo") is None,
    reason="仅显式授权 CN_PATENT_RUN_VISUAL_TESTS=1 时使用 LibreOffice/pdfinfo 导出合成 DOCX",
)
def test_linux_omml_roundtrip_renders_to_pdf(tmp_path):
    module = ASSEMBLER
    document = Document()
    paragraph = document.add_paragraph("公式：")
    registry = module.MathRegistry()
    module.append_omath(paragraph, registry, "C(P,a_i)+(x^2)/(y_1-z)")
    output = tmp_path / "native-math.docx"
    pdf = tmp_path / "native-math.pdf"
    document.save(output)
    module.export_pdf_with_libreoffice(output, pdf)
    assert pdf.is_file() and pdf.stat().st_size > 0
    assert module.count_pdf_pages(pdf) == 1


def test_linux_build_document_generates_native_math_without_word(monkeypatch, tmp_path):
    module = ASSEMBLER
    monkeypatch.setattr(module.sys, "platform", "linux")
    template = tmp_path / "输出模版.docx"
    _write_test_template(template)
    source = tmp_path / "02-申请文件"
    figures = source / "说明书附图"
    figures.mkdir(parents=True)
    (source / "权利要求书.md").write_text(
        "# 权利要求书\n\n1. 一种装置，包括参数 $a_i$。\n", encoding="utf-8"
    )
    (source / "说明书.md").write_text(
        "# 示例装置\n\n## 技术领域\n\n涉及参数 $a_i$。\n\n"
        "## 背景技术\n\n现有方案不足。\n\n## 发明内容\n\n"
        "$$\nC(P,a)=C(P,a_i)+(a-a_i)×[C(P,a_(i+1))-C(P,a_i)]/(a_(i+1)-a_i)\n$$\n\n"
        "## 附图说明\n\n图1为本申请一实施方式中的结构图。\n\n图中：100-示例模块。\n\n"
        "## 具体实施方式\n\n以下结合附图说明本发明的具体实施例。\n\n### 实施例1。\n\n"
        "如图1所示，示例模块100执行计算。\n",
        encoding="utf-8",
    )
    (source / "说明书摘要.md").write_text(
        "# 示例装置\n\n本发明提供一种示例装置。\n\n摘要附图：图1。\n",
        encoding="utf-8",
    )
    Image.new("RGB", (200, 100), "white").save(figures / "图1.png")
    (source / "说明书附图.md").write_text(
        "# 说明书附图\n\n## 图1 结构图\n\n![图1](说明书附图/图1.png)\n",
        encoding="utf-8",
    )
    output = tmp_path / "申请文件.docx"
    report = module.build_document(source, template, output, None)
    assert report["status"] == "STRUCTURE_VERIFIED"
    assert report["math_generation"] == {
        "engine": "direct_omml",
        "platform": "linux",
        "word_automation_required": False,
    }
    assert report["counts"]["native_word_math_objects"] == 3
    assert report["counts"]["pages"] is None
    with ZipFile(output) as archive:
        xml = archive.read("word/document.xml")
    assert xml.count(b"<m:oMath>") == 3
    assert b"<m:f>" in xml
    assert b"<m:sSub>" in xml


def test_visual_review_is_opt_in(monkeypatch, tmp_path):
    module = ASSEMBLER
    monkeypatch.setattr(
        module.sys,
        "argv",
        ["assemble_application_docx.py", "--case-dir", str(tmp_path)],
    )
    args = module.parse_args()
    assert args.visual_review is False

    monkeypatch.setattr(
        module.sys,
        "argv",
        [
            "assemble_application_docx.py",
            "--case-dir",
            str(tmp_path),
            "--visual-review",
        ],
    )
    args = module.parse_args()
    assert args.visual_review is True


# -- F08: 单级模板与多级模板的 numId/ilvl 解析与校验 --


def _template_with_numbering(tmp_path: Path, numbering_xml: str) -> Path:
    """构造一份仅含 ``word/numbering.xml`` 的最小 zip 用于解析层级的单元测试。"""

    path = tmp_path / "模板.docx"
    with ZipFile(path, "w") as archive:
        archive.writestr("word/numbering.xml", numbering_xml)
    return path


def _numbering_xml(entries: list[tuple[int, int, list[tuple[int, list[int]]]]]) -> str:
    """构造 ``word/numbering.xml``。

    ``entries`` 形如 ``[(num_id, abstract_id, [(ilvl, []), ...]), ...]``。
    """

    parts = [
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>',
        '<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">',
    ]
    seen_abstracts: set[int] = set()
    for _num_id, abstract_id, levels in entries:
        if abstract_id not in seen_abstracts:
            parts.append(f'<w:abstractNum w:abstractNumId="{abstract_id}">')
            for ilvl, _ in levels:
                parts.append(
                    f'<w:lvl w:ilvl="{ilvl}"><w:start w:val="1"/></w:lvl>'
                )
            parts.append("</w:abstractNum>")
            seen_abstracts.add(abstract_id)
    for num_id, abstract_id, _ in entries:
        parts.append(
            f'<w:num w:numId="{num_id}">'
            f'<w:abstractNumId w:val="{abstract_id}"/>'
            "</w:num>"
        )
    parts.append("</w:numbering>")
    return "".join(parts)


def test_select_step_ilvl_prefers_sublevel_for_multilevel_template():
    module = ASSEMBLER
    assert module.select_step_ilvl({0, 1, 2}) == 1
    assert module.select_step_ilvl({0, 1, 2, 3}) == 1


def test_select_step_ilvl_falls_back_to_zero_for_single_level_template():
    module = ASSEMBLER
    assert module.select_step_ilvl({0}) == 0
    with pytest.raises(ValueError, match="无法生成步骤段落"):
        module.select_step_ilvl(set())


def test_template_step_ilvl_prefers_demonstrated_level_over_smallest_sublevel():
    module = ASSEMBLER
    document = Document()
    paragraph = document.add_paragraph("权利要求")
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), "29")
    numpr.extend((ilvl, numid))
    paragraph._p.get_or_add_pPr().append(numpr)
    step = document.add_paragraph("S101：步骤")
    step_numpr = OxmlElement("w:numPr")
    step_ilvl = OxmlElement("w:ilvl")
    step_ilvl.set(qn("w:val"), "2")
    step_numid = OxmlElement("w:numId")
    step_numid.set(qn("w:val"), "29")
    step_numpr.extend((step_ilvl, step_numid))
    step._p.get_or_add_pPr().append(step_numpr)

    assert module.template_step_ilvl(document, 29, {0, 1, 2}) == 2
    assert module.template_step_ilvl(Document(), 29, {0, 1, 2}) == 1


def test_template_step_ilvl_rejects_multiple_demonstrated_levels():
    module = ASSEMBLER
    document = Document()
    for level in (1, 2):
        paragraph = document.add_paragraph(f"步骤{level}")
        numpr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(level))
        numid = OxmlElement("w:numId")
        numid.set(qn("w:val"), "29")
        numpr.extend((ilvl, numid))
        paragraph._p.get_or_add_pPr().append(numpr)
    with pytest.raises(ValueError, match="多个步骤层级"):
        module.template_step_ilvl(document, 29, {0, 1, 2})


def test_numbering_style_link_uses_linked_levels_and_rejects_cycles():
    numbering = ASSEMBLER._numbering
    styles = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:styles xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:style w:type="numbering" w:styleId="a">
    <w:name w:val="权利要求书的特殊列表样式"/>
    <w:pPr><w:numPr><w:numId w:val="15"/></w:numPr></w:pPr>
  </w:style>
</w:styles>
""".encode()
    linked = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:abstractNum w:abstractNumId="19">
    <w:lvl w:ilvl="0"><w:start w:val="1"/></w:lvl>
    <w:lvl w:ilvl="2"><w:start w:val="1"/></w:lvl>
  </w:abstractNum>
  <w:abstractNum w:abstractNumId="21">
    <w:numStyleLink w:val="a"/>
  </w:abstractNum>
  <w:num w:numId="15"><w:abstractNumId w:val="19"/></w:num>
  <w:num w:numId="29"><w:abstractNumId w:val="21"/></w:num>
</w:numbering>
""".encode()
    assert numbering._parse_final_numbering(linked, styles) == {15: {0, 2}, 29: {0, 2}}
    with pytest.raises(ValueError, match="缺少 word/styles.xml"):
        numbering._parse_final_numbering(linked, None)

    cycle = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:abstractNum w:abstractNumId="1"><w:numStyleLink w:val="a"/></w:abstractNum>
  <w:num w:numId="15"><w:abstractNumId w:val="1"/></w:num>
</w:numbering>
""".encode()
    with pytest.raises(ValueError, match="编号样式链接循环"):
        numbering._parse_final_numbering(cycle, styles)


def test_repository_template_is_default_and_assembles_claim_steps(tmp_path):
    module = ASSEMBLER
    template = module.project_application_template()
    assert template == ROOT / "模版.docx"
    assert template.is_file()
    assert module.resolve_application_template(None) == template.resolve()
    custom = tmp_path / "其他模板.docx"
    custom.write_bytes(b"PK")
    assert module.resolve_application_template(custom) == custom.resolve()

    source = tmp_path / "02-申请文件"
    figures = source / "说明书附图"
    figures.mkdir(parents=True)
    (source / "权利要求书.md").write_text(
        "# 权利要求书\n\n"
        "1. 一种方法，其特征在于，包括：S101：获取输入；S102：输出结果。\n\n"
        "2. 根据权利要求1所述的方法，其特征在于，还包括校验。\n",
        encoding="utf-8",
    )
    (source / "说明书.md").write_text(
        "# 示例装置\n\n## 技术领域\n\n涉及示例技术。\n\n"
        "## 背景技术\n\n现有方案不足。\n\n## 发明内容\n\n提供示例方案。\n\n"
        "## 附图说明\n\n图1为本申请一实施方式中的结构图。\n\n图中：100-示例模块。\n\n"
        "## 具体实施方式\n\n以下结合附图说明本发明的具体实施例。\n\n### 实施例1。\n\n"
        "如图1所示，示例模块100执行处理。\n",
        encoding="utf-8",
    )
    (source / "说明书摘要.md").write_text(
        "# 示例装置\n\n本发明提供一种示例装置。\n\n摘要附图：图1。\n",
        encoding="utf-8",
    )
    Image.new("RGB", (200, 100), "white").save(figures / "图1.png")
    (source / "说明书附图.md").write_text(
        "# 说明书附图\n\n## 图1 结构图\n\n![图1](说明书附图/图1.png)\n",
        encoding="utf-8",
    )
    output = tmp_path / "申请文件.docx"
    report = module.build_document(source, template, output, None)
    assert report["status"] == "STRUCTURE_VERIFIED"
    assert report["inputs"]["template"] == str(template.resolve())
    errors = []
    module._numbering.check_docx_package(output, errors)
    assert errors == [], errors
    with ZipFile(output) as archive:
        root = etree.fromstring(archive.read("word/document.xml"))
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    steps = []
    for paragraph in root.xpath(".//w:p", namespaces=ns):
        text = "".join(paragraph.xpath(".//w:t/text()", namespaces=ns))
        if not text.startswith("S101"):
            continue
        numid = paragraph.xpath("./w:pPr/w:numPr/w:numId/@w:val", namespaces=ns)
        ilvl = paragraph.xpath("./w:pPr/w:numPr/w:ilvl/@w:val", namespaces=ns)
        steps.append((numid, ilvl, text))
    assert steps == [(["29"], ["2"], "S101：获取输入；")]
    placeholders = (
        "通过skill写专利",
        "泡一杯茶",
        "零度可乐",
        "异构多源医学影像",
        "大模型获取模块",
    )
    with ZipFile(template) as archive:
        template_images = {
            archive.read(name)
            for name in archive.namelist()
            if name.startswith("word/media/")
        }
    with ZipFile(output) as archive:
        visible = []
        for name in archive.namelist():
            if not (name == "word/document.xml" or name.startswith("word/header") or name.startswith("word/footer")):
                continue
            visible.append(archive.read(name).decode("utf-8"))
        visible_text = "\n".join(visible)
        for phrase in placeholders:
            assert phrase not in visible_text
        embeds = set(etree.fromstring(archive.read("word/document.xml")).xpath(
            ".//*[local-name()='blip']/@*[local-name()='embed']"
        ))
        rels = etree.fromstring(archive.read("word/_rels/document.xml.rels"))
        targets = {
            rel.get("Id"): rel.get("Target")
            for rel in rels
        }
        for embed in embeds:
            target = "word/" + targets[embed].lstrip("/")
            assert archive.read(target) not in template_images


def test_missing_project_template_does_not_use_case_template(tmp_path, monkeypatch):
    module = ASSEMBLER
    monkeypatch.setattr(module, "project_application_template", lambda: tmp_path / "模版.docx")
    (tmp_path / "输出模版.docx").write_bytes(b"PK")
    with pytest.raises(FileNotFoundError, match="输出模版.docx"):
        module.resolve_application_template(None)


def test_claim_step_properties_picks_legal_ilvl_per_template(tmp_path):
    module = ASSEMBLER
    document = Document()
    paragraph = document.add_paragraph()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    numid = OxmlElement("w:numId")
    numid.set(qn("w:val"), "1")
    numpr.extend((ilvl, numid))
    paragraph._p.get_or_add_pPr().append(numpr)

    single = module.claim_step_properties(paragraph._p.pPr, {0})
    single_ilvl = single.find(qn("w:numPr")).find(qn("w:ilvl"))
    assert single_ilvl.get(qn("w:val")) == "0"

    multilevel = module.claim_step_properties(paragraph._p.pPr, {0, 1, 2})
    multi_ilvl = multilevel.find(qn("w:numPr")).find(qn("w:ilvl"))
    assert multi_ilvl.get(qn("w:val")) == "1"


def test_parse_template_numbering_collects_abstract_and_override(tmp_path):
    module = ASSEMBLER
    # numId=1 绑定 abstractNum=0（含 ilvl 0/1/2）；numId=2 通过 lvlOverride 重写为仅 ilvl 0/2。
    numbering_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:abstractNum w:abstractNumId="0">
    <w:lvl w:ilvl="0"><w:start w:val="1"/></w:lvl>
    <w:lvl w:ilvl="1"><w:start w:val="1"/></w:lvl>
    <w:lvl w:ilvl="2"><w:start w:val="1"/></w:lvl>
  </w:abstractNum>
  <w:abstractNum w:abstractNumId="1">
    <w:lvl w:ilvl="0"><w:start w:val="1"/></w:lvl>
  </w:abstractNum>
  <w:num w:numId="1">
    <w:abstractNumId w:val="0"/>
  </w:num>
  <w:num w:numId="2">
    <w:abstractNumId w:val="1"/>
    <w:lvlOverride w:ilvl="2"><w:lvl w:ilvl="2"><w:start w:val="1"/></w:lvl></w:lvlOverride>
  </w:num>
</w:numbering>
"""
    template = _template_with_numbering(tmp_path, numbering_xml)
    parsed = module.parse_template_numbering(template)
    # 真实 w:lvl 为第二个编号实例增加 ilvl=2；保留基底 ilvl=0。
    assert parsed == {1: {0, 1, 2}, 2: {0, 2}}


def test_parse_template_numbering_handles_corrupt_numbering_xml(tmp_path):
    module = ASSEMBLER
    template = tmp_path / "坏numbering.docx"
    with ZipFile(template, "w") as archive:
        archive.writestr("word/numbering.xml", "<not-xml>")
    with pytest.raises(ValueError, match="numbering.xml"):
        module.parse_template_numbering(template)


def test_parse_template_numbering_override_requires_real_lvl(tmp_path):
    module = ASSEMBLER
    numbering_xml = """<w:numbering xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0"/></w:abstractNum>
      <w:num w:numId="1"><w:abstractNumId w:val="0"/><w:lvlOverride w:ilvl="2"><w:lvl w:ilvl="2"/></w:lvlOverride></w:num>
    </w:numbering>"""
    template = _template_with_numbering(tmp_path, numbering_xml)
    assert module.parse_template_numbering(template) == {1: {0, 2}}


def test_parse_template_numbering_missing_member_is_unavailable(tmp_path):
    module = ASSEMBLER
    template = tmp_path / "缺失numbering.docx"
    with ZipFile(template, "w") as archive:
        archive.writestr("word/document.xml", "<w:document xmlns:w=\"http://schemas.openxmlformats.org/wordprocessingml/2006/main\"/>")
    assert module.parse_template_numbering(template) is None


def test_parse_template_numbering_python_docx_default_template_is_single_level(tmp_path):
    """python-docx 默认模板自带 ``numbering.xml``，每个 numId 只声明 ``ilvl=0``。"""

    module = ASSEMBLER
    template = tmp_path / "默认模板.docx"
    _write_test_template(template)
    parsed = module.parse_template_numbering(template)
    assert parsed, "python-docx 默认模板应提供 numbering.xml"
    # 默认模板所有 numId 均为单级列表（仅 ilvl=0），与单级步骤段落合同一致。
    for levels in parsed.values():
        assert levels == {0}


def test_validate_numbering_references_rejects_undefined_numid():
    module = ASSEMBLER
    document_xml = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b'<w:body><w:p><w:pPr><w:numPr><w:ilvl w:val="0"/>'
        b'<w:numId w:val="99"/></w:numPr></w:pPr></w:p></w:body></w:document>'
    )
    with pytest.raises(ValueError, match="未定义的 numId=99"):
        module.validate_numbering_references(document_xml, {1: {0}})


def test_validate_numbering_references_rejects_undefined_ilvl():
    module = ASSEMBLER
    document_xml = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b'<w:body><w:p><w:pPr><w:numPr><w:ilvl w:val="2"/>'
        b'<w:numId w:val="1"/></w:numPr></w:pPr></w:p></w:body></w:document>'
    )
    with pytest.raises(ValueError, match="未定义层级 ilvl=2"):
        module.validate_numbering_references(document_xml, {1: {0}})


def test_validate_numbering_references_accepts_legal_ilvl():
    module = ASSEMBLER
    document_xml = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b'<w:body>'
        b'<w:p><w:pPr><w:numPr><w:ilvl w:val="0"/>'
        b'<w:numId w:val="1"/></w:numPr></w:pPr></w:p>'
        b'<w:p><w:pPr><w:numPr><w:ilvl w:val="1"/>'
        b'<w:numId w:val="1"/></w:numPr></w:pPr></w:p>'
        b'</w:body></w:document>'
    )
    module.validate_numbering_references(document_xml, {1: {0, 1, 2}})


def test_validate_numbering_references_rejects_missing_numbering_when_referenced():
    module = ASSEMBLER
    document_xml = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b'<w:body><w:p><w:pPr><w:numPr><w:ilvl w:val="2"/>'
        b'<w:numId w:val="1"/></w:numPr></w:pPr></w:p></w:body></w:document>'
    )
    with pytest.raises(ValueError, match="缺少或损坏"):
        module.validate_numbering_references(document_xml, None)


def test_validate_numbering_references_allows_missing_numbering_without_references():
    module = ASSEMBLER
    document_xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        '<w:body><w:p><w:r><w:t>无编号正文</w:t></w:r></w:p></w:body></w:document>'
    ).encode('utf-8')
    module.validate_numbering_references(document_xml, None)


def test_validate_numbering_references_accepts_numid_zero_as_no_numbering():
    module = ASSEMBLER
    document_xml = (
        b'<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
        b'<w:body><w:p><w:pPr><w:numPr><w:ilvl w:val="0"/>'
        b'<w:numId w:val="0"/></w:numPr></w:pPr></w:p></w:body></w:document>'
    )
    module.validate_numbering_references(document_xml, None)


def test_linux_build_document_with_single_level_template_does_not_reference_undefined_ilvl(
    monkeypatch, tmp_path
):
    """F08 正例：模板只声明 ``ilvl=0``，步骤段落仍以 ``ilvl=0`` 输出，验证通过。"""

    module = ASSEMBLER
    monkeypatch.setattr(module.sys, "platform", "linux")
    template = tmp_path / "输出模版.docx"
    _write_test_template(template)
    source = tmp_path / "02-申请文件"
    figures = source / "说明书附图"
    figures.mkdir(parents=True)
    (source / "权利要求书.md").write_text(
        "# 权利要求书\n\n"
        "1. 一种方法，其特征在于，包括：S101：执行第一步骤；"
        "S102：执行第二步骤；S103：执行第三步骤。\n",
        encoding="utf-8",
    )
    (source / "说明书.md").write_text(
        "# 示例装置\n\n## 技术领域\n\n涉及示例。\n\n"
        "## 背景技术\n\n现有方案不足。\n\n## 发明内容\n\n提供示例方案。\n\n"
        "## 附图说明\n\n图1为本申请一实施方式中的结构图。\n\n图中：100-示例模块。\n\n"
        "## 具体实施方式\n\n以下结合附图说明本发明的具体实施例。\n\n### 实施例1。\n\n"
        "如图1所示，示例模块100执行处理。\n",
        encoding="utf-8",
    )
    (source / "说明书摘要.md").write_text(
        "# 示例装置\n\n本发明提供一种示例装置。\n\n摘要附图：图1。\n",
        encoding="utf-8",
    )
    Image.new("RGB", (200, 100), "white").save(figures / "图1.png")
    (source / "说明书附图.md").write_text(
        "# 说明书附图\n\n## 图1 结构图\n\n![图1](说明书附图/图1.png)\n",
        encoding="utf-8",
    )
    output = tmp_path / "申请文件.docx"
    report = module.build_document(source, template, output, None)
    assert report["status"] == "STRUCTURE_VERIFIED"
    with ZipFile(output) as archive:
        document_xml = archive.read("word/document.xml")
    parsed = etree.fromstring(document_xml)
    ilvl_values = [
        ilvl.get(qn("w:val"))
        for ilvl in parsed.iter(qn("w:ilvl"))
    ]
    # 单级模板：步骤段落不应出现 ilvl=2，全部段落均使用模板已声明的 ilvl=0。
    assert "2" not in ilvl_values
    assert all(value == "0" for value in ilvl_values)


def test_normalize_formula_with_latex_symbols():
    """测试 LaTeX 符号宏转 Unicode：最长匹配、函数名、不支持的宏报错。"""
    module = ASSEMBLER

    # 1. 产品符号与运算符正确处理
    assert module.normalize_formula(
        r"\prod_{k \in K_c} I(r_k=a_k) \cdot I(x)"
    ) == "∏_{k ∈ K_c} I(r_k=a_k) · I(x)"

    # 2. 避免 \cdot 截断 \cdots；检验 \int 与 \infty 的独立性
    assert module.normalize_formula(r"a \cdots b") == "a ⋯ b"
    assert module.normalize_formula(r"\int_0^\infty x \in X") == "∫_0^∞ x ∈ X"

    # 3. 函数名宏转为普通标识符
    assert module.normalize_formula(
        r"\alpha_i + \max_{w} f(w)"
    ) == "α_i + max_{w} f(w)"

    # 4. \left \right 忽略、多余空格被去掉
    result = module.normalize_formula(r"\left( a \right)")
    assert result.replace(" ", "") == "(a)"
    # 解析不抛异常
    parsed = module.FormulaParser(result).parse()
    assert parsed is not None

    # 5. 转义字符 \{ \} 转为字面字符，可解析
    assert module.normalize_formula(r"\{a,b\}") == "{a,b}"
    parsed = module.FormulaParser("{a,b}").parse()
    assert parsed.kind == "delimiter" and parsed.value == "{}"

    # 6. 不支持的宏报错
    with pytest.raises(ValueError) as exc_info:
        module.normalize_formula(r"\frac{a}{b}")
    assert r"\frac" in str(exc_info.value)

    with pytest.raises(ValueError) as exc_info:
        module.normalize_formula(r"TTL\_inv")
    assert r"\_" in str(exc_info.value)

    with pytest.raises(ValueError) as exc_info:
        module.normalize_formula(r"\hat{y}")
    assert r"\hat" in str(exc_info.value)


def test_normalize_formula_build_omath_with_latex_symbols(tmp_path):
    """测试 LaTeX 符号在 OMML 中的正确结构和无反斜杠。"""
    module = ASSEMBLER

    # 用 \prod 和 \in 构建公式，验证 OMML 输出
    formula = r"\prod_{k \in K} x_k"
    root = module.build_omath(formula)
    xml = etree.tostring(root)

    # 断言不含反斜杠
    assert b"\\" not in xml

    # 断言 XML 中含下标结构
    namespaces = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}
    parsed = etree.fromstring(xml)
    assert len(parsed.xpath(".//m:sSub", namespaces=namespaces)) >= 1


def test_normalize_formula_legacy_block_formats():
    """验证现有的块格式处理不回归。"""
    module = ASSEMBLER

    assert module.normalize_formula("$$ a_i + b $$") == "a_i + b"
    assert module.normalize_formula("```math\nx = y\n```") == "x = y"

    # 验证尾随 。 被去掉
    assert module.normalize_formula("$$ a + b 。 $$") == "a + b"


def test_pending_decisions_markers(tmp_path):
    import subprocess
    import json
    from zipfile import ZipFile

    case_dir = tmp_path / "case"
    source_dir = case_dir / "02-申请文件"
    source_dir.mkdir(parents=True)
    template_path = case_dir / "输出模版.docx"
    _write_test_template(template_path)

    (source_dir / "说明书.md").write_text("# 测试发明\n\n## 技术领域\n\n测试领域【待决-D002】\n\n## 背景技术\n\n测试背景\n\n## 发明内容\n\n内容\n\n## 附图说明\n\n图1为结构示意图。\n\n图中：1。 \n\n## 具体实施方式\n\n以下结合附图说明本发明的具体实施例。\n\n### 实施例1。\n\n结合图1所示，具体测试【待决-D003】\n", "utf-8")
    (source_dir / "权利要求书.md").write_text("# 权利要求书\n\n1. 一种测试。\n2. 如权利要求1的测试【待决-D001】\n", "utf-8")
    (source_dir / "说明书摘要.md").write_text("# 说明书摘要\n\n摘要正文\n\n摘要附图：图1。\n", "utf-8")
    (source_dir / "说明书附图.md").write_text("## 图1 结构示意图\n![](fig1.png)\n", "utf-8")
    Image.new("RGB", (100, 100)).save(source_dir / "fig1.png")

    script_path = ROOT / "skills/cn-patent-application-creator/scripts/assemble_application_docx.py"
    verify_script = ROOT / "skills/cn-patent-application-creator/scripts/verify_docx_assembly.py"

    # review copy
    out_review = case_dir / "review.docx"
    cmd_review = [sys.executable, str(script_path), "--case-dir", str(case_dir), "--output", str(out_review), "--copy", "review"]
    res = subprocess.run(cmd_review, capture_output=True, text=True)
    assert res.returncode == 0

    with ZipFile(out_review) as z:
        xml = z.read("word/document.xml").decode("utf-8")
        assert "【待决-D001】" in xml
        assert "【待决-D002】" in xml
        assert 'w:highlight' in xml
        assert 'yellow' in xml

    work_dirs = sorted((case_dir / "03-审查工作区").glob("docx组装-*"))
    review_report_dir = work_dirs[-1]
    review_report = review_report_dir / "docx-assembly-report.json"

    res_verify = subprocess.run([sys.executable, str(verify_script), "--report", str(review_report)], capture_output=True)
    assert res_verify.returncode == 0

    # submission copy
    out_sub = case_dir / "submission.docx"
    cmd_sub = [sys.executable, str(script_path), "--case-dir", str(case_dir), "--output", str(out_sub), "--copy", "submission"]
    res_sub = subprocess.run(cmd_sub, capture_output=True, text=True)
    assert res_sub.returncode == 0

    with ZipFile(out_sub) as z:
        xml = z.read("word/document.xml").decode("utf-8")
        assert "【待决-D001】" not in xml
        assert "【待决-D002】" not in xml
        assert "【待决-D003】" not in xml
        assert 'w:highlight w:val="yellow"' not in xml

    work_dirs = sorted((case_dir / "03-审查工作区").glob("docx组装-*"))
    sub_report_dir = work_dirs[-1]
    sub_report = sub_report_dir / "docx-assembly-report.json"

    res_verify_sub = subprocess.run([sys.executable, str(verify_script), "--report", str(sub_report)], capture_output=True)
    assert res_verify_sub.returncode == 0

    # broken marker submission
    (source_dir / "权利要求书.md").write_text("# 权利要求书\n\n1. 一种测试。\n2. 如权利要求1的测试【待决-D0】\n", "utf-8")
    out_broken = case_dir / "broken.docx"
    cmd_broken = [sys.executable, str(script_path), "--case-dir", str(case_dir), "--output", str(out_broken), "--copy", "submission"]
    res_broken = subprocess.run(cmd_broken, capture_output=True, text=True)
    assert res_broken.returncode != 0
    assert "提交副本存在残缺待决标记" in res_broken.stderr
